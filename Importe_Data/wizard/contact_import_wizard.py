# -*- coding: utf-8 -*-
"""Asistente para importar contactos desde archivos CSV/XML/Excel."""

import base64
import csv
import io
import logging
import xml.etree.ElementTree as ET
from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    _logger.warning("openpyxl no está instalado. La importación de Excel no estará disponible.")


class ContactImportConflict(models.TransientModel):
    """Conflicto detectado durante la importación."""
    
    _name = 'contact.import.conflict'
    _description = 'Conflicto de Importación'
    _order = 'row_number'

    wizard_id = fields.Many2one('contact.import.wizard', string='Wizard', required=True, ondelete='cascade')
    row_number = fields.Integer(string='Fila #', required=True)
    
    # Datos del archivo
    new_name = fields.Char(string='Nombre (Archivo)')
    new_vat = fields.Char(string='Cédula (Archivo)')
    new_email = fields.Char(string='Email (Archivo)')
    new_phone = fields.Char(string='Teléfono (Archivo)')
    new_mobile = fields.Char(string='Móvil (Archivo)')
    
    # Contacto existente
    existing_partner_id = fields.Many2one('res.partner', string='Contacto Existente')
    existing_name = fields.Char(string='Nombre (Sistema)', related='existing_partner_id.name', readonly=True)
    existing_vat = fields.Char(string='Cédula (Sistema)', related='existing_partner_id.vat', readonly=True)
    existing_email = fields.Char(string='Email (Sistema)', related='existing_partner_id.email', readonly=True)
    
    # Tipo de conflicto
    conflict_type = fields.Selection([
        ('vat', 'Cédula duplicada'),
        ('email', 'Email duplicado'),
        ('name', 'Nombre similar'),
        ('multiple', 'Múltiples coincidencias'),
    ], string='Tipo de Conflicto', required=True)
    
    # Acción a tomar
    action = fields.Selection([
        ('update', 'Actualizar existente'),
        ('create', 'Crear nuevo'),
        ('skip', 'Omitir esta fila'),
    ], string='Acción', default='update', required=True)
    
    # Datos temporales para la importación
    import_data = fields.Text(string='Datos de Importación', help='JSON con los datos a importar')


class ContactImportColumn(models.TransientModel):
    """Columna de importación con mapeo a campos de Odoo."""
    
    _name = 'contact.import.column'
    _description = 'Columna de Importación de Contactos'
    _order = 'column_index'

    wizard_id = fields.Many2one('contact.import.wizard', string='Wizard', required=True, ondelete='cascade')
    column_name = fields.Char(string='Nombre de Columna', required=True)
    column_index = fields.Integer(string='Índice', required=True)
    
    odoo_field = fields.Selection([
        ('name', 'Nombre'),
        ('vat', 'Documento/NIF'),
        ('email', 'Email'),
        ('phone', 'Teléfono'),
        ('mobile', 'Móvil'),
        ('street', 'Dirección'),
        ('city', 'Ciudad'),
        ('country_id', 'País'),
        ('comment', 'Notas/Grado/Curso'),
    ], string='Campo de Odoo')
    
    odoo_field_label = fields.Char(string='Etiqueta del Campo', compute='_compute_odoo_field_label')

    @api.depends('odoo_field')
    def _compute_odoo_field_label(self):
        """Calcula la etiqueta legible del campo de Odoo."""
        for record in self:
            if record.odoo_field:
                selection_dict = dict(self._fields['odoo_field'].selection)
                record.odoo_field_label = selection_dict.get(record.odoo_field, '')
            else:
                record.odoo_field_label = ''


class ContactImportWizard(models.TransientModel):
    """Asistente para importar contactos desde archivo."""
    
    _name = 'contact.import.wizard'
    _description = 'Asistente de Importación de Contactos'

    # ========== PASO 1: Subir archivo ==========
    step = fields.Selection([
        ('upload', 'Cargar Archivo'),
        ('mapping', 'Mapear Campos'),
        ('preview', 'Previsualizar'),
        ('conflicts', 'Resolver Conflictos'),
        ('confirm', 'Confirmar Importación'),
        ('result', 'Resultado'),
    ], string='Paso', default='upload', required=True)
    
    file_data = fields.Binary(string='Archivo', required=True, attachment=False)
    file_name = fields.Char(string='Nombre de Archivo')
    file_type = fields.Selection([
        ('csv', 'CSV'),
        ('xlsx', 'Excel (XLSX)'),
        ('xml', 'XML'),
    ], string='Tipo de Archivo', compute='_compute_file_type', store=True)
    
    # ========== PASO 2: Configuración ==========
    contact_type = fields.Selection([
        ('student', 'Estudiante'),
        ('teacher', 'Profesor'),
    ], string='Tipo de Contacto', default='student', required=True)
    
    region = fields.Char(string='Región', help='Región de origen de los contactos')
    
    # CSV específico
    csv_delimiter = fields.Selection([
        (',', 'Coma (,)'),
        (';', 'Punto y coma (;)'),
        ('\t', 'Tabulador'),
        ('|', 'Tubería (|)'),
    ], string='Delimitador CSV', default=',')
    
    csv_encoding = fields.Selection([
        ('utf-8', 'UTF-8'),
        ('latin-1', 'Latin-1'),
        ('cp1252', 'Windows-1252'),
    ], string='Codificación', default='utf-8')
    
    has_header = fields.Boolean(string='Primera fila es encabezado', default=True)
    
    # ========== PASO 3: Mapeo de campos ==========
    column_ids = fields.One2many('contact.import.column', 'wizard_id', string='Columnas')
    preview_data = fields.Html(string='Vista Previa de Datos', readonly=True)
    
    # ========== PASO 4: Conflictos ==========
    conflict_ids = fields.One2many('contact.import.conflict', 'wizard_id', string='Conflictos')
    has_conflicts = fields.Boolean(string='Tiene Conflictos', compute='_compute_has_conflicts', store=True)
    conflict_count = fields.Integer(string='Total Conflictos', compute='_compute_has_conflicts', store=True)
    
    # ========== PASO 5: Resultados ==========
    total_rows = fields.Integer(string='Total de Filas', readonly=True)
    to_create = fields.Integer(string='A Crear', readonly=True)
    to_update = fields.Integer(string='A Actualizar', readonly=True)
    with_errors = fields.Integer(string='Con Errores', readonly=True)
    error_log = fields.Text(string='Log de Errores', readonly=True)
    
    created_ids = fields.Many2many('res.partner', 'contact_wizard_partner_created_rel',
                                   'wizard_id', 'partner_id',
                                   string='Contactos Creados', readonly=True)
    updated_ids = fields.Many2many('res.partner', 'contact_wizard_partner_updated_rel',
                                   'wizard_id', 'partner_id', 
                                   string='Contactos Actualizados', readonly=True)

    @api.depends('file_name')
    def _compute_file_type(self):
        """Detecta el tipo de archivo basado en la extensión."""
        for record in self:
            if record.file_name:
                file_name_lower = record.file_name.lower()
                if file_name_lower.endswith('.csv'):
                    record.file_type = 'csv'
                elif file_name_lower.endswith('.xlsx'):
                    record.file_type = 'xlsx'
                elif file_name_lower.endswith('.xml'):
                    record.file_type = 'xml'
                else:
                    record.file_type = False
            else:
                record.file_type = False

    @api.depends('conflict_ids')
    def _compute_has_conflicts(self):
        """Calcula si hay conflictos y cuántos."""
        for record in self:
            record.conflict_count = len(record.conflict_ids)
            record.has_conflicts = record.conflict_count > 0

    def action_next_step(self):
        """Avanza al siguiente paso del wizard."""
        self.ensure_one()
        
        if self.step == 'upload':
            # Validar archivo y pasar a mapeo
            if not self.file_data or not self.file_type:
                raise UserError(_('Debe cargar un archivo válido (CSV, XLSX o XML).'))
            
            self._parse_file()
            
            if not self.column_ids:
                raise UserError(_('No se pudieron detectar columnas en el archivo.'))
            
            self.step = 'mapping'
            
        elif self.step == 'mapping':
            # Validar mapeo y generar vista previa
            if not any(col.odoo_field for col in self.column_ids):
                raise UserError(_('Debe mapear al menos un campo antes de continuar.'))
            
            self._generate_preview()
            self.step = 'preview'
            
        elif self.step == 'preview':
            # Detectar conflictos antes de importar
            self._detect_conflicts()
            
            if self.has_conflicts:
                self.step = 'conflicts'
            else:
                # Si no hay conflictos, proceder directamente a importar
                self._process_import()
                self.step = 'result'
                
        elif self.step == 'conflicts':
            # Procesar importación con las decisiones tomadas
            self._process_import()
            self.step = 'result'
        
        return self._reopen_wizard()

    def action_previous_step(self):
        """Retrocede al paso anterior."""
        self.ensure_one()
        
        if self.step == 'mapping':
            self.step = 'upload'
        elif self.step == 'preview':
            self.step = 'mapping'
        elif self.step == 'conflicts':
            self.step = 'preview'
        elif self.step == 'result':
            self.step = 'conflicts' if self.has_conflicts else 'preview'
        
        return self._reopen_wizard()

    def _reopen_wizard(self):
        """Reabre el wizard en el mismo registro."""
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _parse_file(self):
        """Lee el archivo y crea las columnas disponibles."""
        self.ensure_one()
        
        # Limpiar columnas previas
        self.column_ids.unlink()
        
        if self.file_type == 'csv':
            columns = self._parse_csv()
        elif self.file_type == 'xlsx':
            columns = self._parse_xlsx()
        elif self.file_type == 'xml':
            columns = self._parse_xml()
        else:
            raise UserError(_('Tipo de archivo no soportado.'))
        
        # Crear registros de columnas directamente
        for idx, col_name in enumerate(columns):
            self.env['contact.import.column'].create({
                'wizard_id': self.id,
                'column_name': col_name,
                'column_index': idx,
                'odoo_field': self._auto_map_field(col_name),
            })
        
        # Forzar actualización del caché
        self.invalidate_recordset(['column_ids'])

    def _parse_csv(self):
        """Extrae columnas de un archivo CSV."""
        try:
            file_content = base64.b64decode(self.file_data)
            
            try:
                decoded = file_content.decode(self.csv_encoding)
            except UnicodeDecodeError:
                decoded = file_content.decode('latin-1')
            
            csv_reader = csv.reader(io.StringIO(decoded), delimiter=self.csv_delimiter)
            
            if self.has_header:
                header = next(csv_reader, None)
                return header if header else []
            else:
                first_row = next(csv_reader, None)
                return [f'Columna {i+1}' for i in range(len(first_row))] if first_row else []
                
        except UnicodeDecodeError:
            raise UserError(_('Error de codificación. Intente con otra codificación (UTF-8, Latin-1, Windows-1252).'))
        except Exception as e:
            raise UserError(_('Error al leer el archivo CSV: %s') % str(e))

    def _parse_xlsx(self):
        """Extrae columnas de un archivo Excel."""
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('La librería openpyxl no está instalada. Instálela con: pip install openpyxl'))
        
        try:
            file_content = base64.b64decode(self.file_data)
            wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
            ws = wb.active
            
            if self.has_header:
                header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                return [str(cell) if cell else f'Columna {i+1}' for i, cell in enumerate(header_row)] if header_row else []
            else:
                first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                return [f'Columna {i+1}' for i in range(len(first_row))] if first_row else []
                
        except Exception as e:
            raise UserError(_('Error al leer el archivo Excel: %s') % str(e))

    def _parse_xml(self):
        """Extrae columnas de un archivo XML."""
        try:
            file_content = base64.b64decode(self.file_data)
            root = ET.fromstring(file_content)
            
            # Buscar el primer registro para obtener los campos
            first_record = root.find('.//record') or root.find('.//row') or root.find('.//contact')
            
            if first_record is not None:
                columns = [child.tag for child in first_record]
                return columns
            else:
                raise UserError(_('No se encontraron registros en el archivo XML.'))
            
        except ET.ParseError:
            raise UserError(_('El archivo XML no tiene un formato válido.'))
        except Exception as e:
            raise UserError(_('Error al leer el archivo XML: %s') % str(e))

    def _auto_map_field(self, column_name):
        """Intenta mapear automáticamente una columna a un campo de Odoo."""
        if not column_name:
            return False
        
        column_lower = column_name.lower().strip()
        
        # Mapeo automático basado en palabras clave
        mappings = {
            'nombre': 'name',
            'name': 'name',
            'apellido': 'name',
            'documento': 'vat',
            'vat': 'vat',
            'dni': 'vat',
            'identificacion': 'vat',
            'cedula': 'vat',
            'email': 'email',
            'correo': 'email',
            'mail': 'email',
            'telefono': 'phone',
            'phone': 'phone',
            'tel': 'phone',
            'celular': 'mobile',
            'mobile': 'mobile',
            'movil': 'mobile',
            'direccion': 'street',
            'street': 'street',
            'calle': 'street',
            'ciudad': 'city',
            'city': 'city',
            'pais': 'country_id',
            'country': 'country_id',
            'grado': 'comment',
            'curso': 'comment',
            'nivel': 'comment',
            'comentario': 'comment',
            'comment': 'comment',
            'notas': 'comment',
        }
        
        for key, field in mappings.items():
            if key in column_lower:
                return field
        
        return False

    def _generate_preview(self):
        """Genera una vista previa de los datos a importar."""
        self.ensure_one()
        
        rows = self._read_data_rows(max_rows=10)
        
        if not rows:
            self.preview_data = '<p class="text-muted">No se encontraron datos para previsualizar.</p>'
            return
        
        # Crear tabla HTML para previsualización
        preview_html = '<table class="table table-sm table-bordered">'
        
        # Encabezado
        preview_html += '<thead><tr>'
        for col in self.column_ids.sorted('column_index'):
            field_label = f'<br/><small class="text-muted">({col.odoo_field_label})</small>' if col.odoo_field else ''
            preview_html += f'<th>{col.column_name}{field_label}</th>'
        preview_html += '</tr></thead>'
        
        # Datos
        preview_html += '<tbody>'
        for row in rows:
            preview_html += '<tr>'
            for value in row:
                preview_html += f'<td>{value or ""}</td>'
            preview_html += '</tr>'
        preview_html += '</tbody>'
        
        preview_html += '</table>'
        
        self.preview_data = preview_html

    def _detect_conflicts(self):
        """Detecta conflictos con contactos existentes."""
        self.ensure_one()
        
        # Limpiar conflictos previos
        self.conflict_ids.unlink()
        
        rows = self._read_data_rows()
        if not rows:
            return
        
        # Obtener mapeo de columnas
        field_mapping = {
            col.column_index: col.odoo_field 
            for col in self.column_ids 
            if col.odoo_field
        }
        
        conflicts = []
        
        for row_idx, row in enumerate(rows, start=1):
            try:
                # Extraer valores
                values = self._extract_values_from_row(row, field_mapping)
                
                if not values:
                    continue
                
                # Buscar conflictos por cédula, email o nombre
                conflict_data = self._find_conflicts(values, row_idx)
                
                if conflict_data:
                    conflicts.append(conflict_data)
                    
            except Exception as e:
                _logger.warning(f'Error al detectar conflictos en fila {row_idx}: {e}')
        
        # Crear registros de conflictos
        for conflict in conflicts:
            self.env['contact.import.conflict'].create(conflict)

    def _find_conflicts(self, values, row_number):
        """Busca conflictos con contactos existentes basándose en cédula, email o nombre."""
        import json
        
        vat = values.get('vat', '').strip() if values.get('vat') else ''
        email = values.get('email', '').strip() if values.get('email') else ''
        name = values.get('name', '').strip() if values.get('name') else ''
        
        if not vat and not email and not name:
            return None
        
        # Buscar por cédula (prioridad más alta)
        if vat:
            partner = self.env['res.partner'].search([('vat', '=ilike', vat)], limit=1)
            if partner:
                return {
                    'wizard_id': self.id,
                    'row_number': row_number,
                    'new_name': name,
                    'new_vat': vat,
                    'new_email': email,
                    'new_phone': values.get('phone', ''),
                    'new_mobile': values.get('mobile', ''),
                    'existing_partner_id': partner.id,
                    'conflict_type': 'vat',
                    'action': 'update',
                    'import_data': json.dumps(values),
                }
        
        # Buscar por email
        if email:
            partner = self.env['res.partner'].search([('email', '=ilike', email)], limit=1)
            if partner:
                return {
                    'wizard_id': self.id,
                    'row_number': row_number,
                    'new_name': name,
                    'new_vat': vat,
                    'new_email': email,
                    'new_phone': values.get('phone', ''),
                    'new_mobile': values.get('mobile', ''),
                    'existing_partner_id': partner.id,
                    'conflict_type': 'email',
                    'action': 'update',
                    'import_data': json.dumps(values),
                }
        
        # Buscar por nombre exacto
        if name:
            partner = self.env['res.partner'].search([('name', '=ilike', name)], limit=1)
            if partner:
                return {
                    'wizard_id': self.id,
                    'row_number': row_number,
                    'new_name': name,
                    'new_vat': vat,
                    'new_email': email,
                    'new_phone': values.get('phone', ''),
                    'new_mobile': values.get('mobile', ''),
                    'existing_partner_id': partner.id,
                    'conflict_type': 'name',
                    'action': 'update',
                    'import_data': json.dumps(values),
                }
        
        return None

    def _read_data_rows(self, max_rows=None):
        """Lee las filas de datos del archivo."""
        self.ensure_one()
        
        if self.file_type == 'csv':
            return self._read_csv_rows(max_rows)
        elif self.file_type == 'xlsx':
            return self._read_xlsx_rows(max_rows)
        elif self.file_type == 'xml':
            return self._read_xml_rows(max_rows)
        
        return []

    def _read_csv_rows(self, max_rows=None):
        """Lee filas de datos de un CSV."""
        try:
            file_content = base64.b64decode(self.file_data)
            decoded = file_content.decode(self.csv_encoding)
            csv_reader = csv.reader(io.StringIO(decoded), delimiter=self.csv_delimiter)
            
            rows = []
            start_row = 1 if self.has_header else 0
            
            for idx, row in enumerate(csv_reader):
                if idx < start_row:
                    continue
                rows.append(row)
                if max_rows and len(rows) >= max_rows:
                    break
            
            return rows
            
        except Exception as e:
            raise UserError(_('Error al leer datos CSV: %s') % str(e))

    def _read_xlsx_rows(self, max_rows=None):
        """Lee filas de datos de un Excel."""
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('openpyxl no está disponible.'))
        
        try:
            file_content = base64.b64decode(self.file_data)
            wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
            ws = wb.active
            
            start_row = 2 if self.has_header else 1
            rows = []
            
            for row in ws.iter_rows(min_row=start_row, values_only=True):
                rows.append([str(cell) if cell is not None else '' for cell in row])
                if max_rows and len(rows) >= max_rows:
                    break
            
            return rows
            
        except Exception as e:
            raise UserError(_('Error al leer datos Excel: %s') % str(e))

    def _read_xml_rows(self, max_rows=None):
        """Lee filas de datos de un XML."""
        try:
            file_content = base64.b64decode(self.file_data)
            root = ET.fromstring(file_content)
            
            records = root.findall('.//record') or root.findall('.//row') or root.findall('.//contact')
            
            rows = []
            column_order = [col.column_name for col in self.column_ids.sorted('column_index')]
            
            for record in records:
                row = []
                for col_name in column_order:
                    child = record.find(col_name)
                    row.append(child.text if child is not None and child.text else '')
                rows.append(row)
                
                if max_rows and len(rows) >= max_rows:
                    break
            
            return rows
            
        except Exception as e:
            raise UserError(_('Error al leer datos XML: %s') % str(e))

    def _process_import(self):
        """Procesa la importación de contactos con las decisiones de conflictos."""
        self.ensure_one()
        
        import json
        
        rows = self._read_data_rows()
        
        if not rows:
            raise UserError(_('No hay datos para importar.'))
        
        # Preparar contadores y logs
        to_create = []
        to_update = []
        errors = []
        created_partners = self.env['res.partner']
        updated_partners = self.env['res.partner']
        
        # Información de trazabilidad
        import_timestamp = fields.Datetime.now()
        import_info = {
            'import_source': self.file_name or 'Archivo sin nombre',
            'import_date': import_timestamp,
            'import_region': self.region or False,
            'import_type': self.contact_type,
            'last_import_update': import_timestamp,
        }
        
        # Obtener mapeo de columnas
        field_mapping = {
            col.column_index: col.odoo_field 
            for col in self.column_ids 
            if col.odoo_field
        }
        
        # Crear diccionario de conflictos por número de fila
        conflicts_by_row = {
            conflict.row_number: conflict 
            for conflict in self.conflict_ids
        }
        
        for row_idx, row in enumerate(rows, start=1):
            try:
                # Extraer valores según el mapeo
                values = self._extract_values_from_row(row, field_mapping)
                
                if not values:
                    errors.append(f'Fila {row_idx}: No se pudieron extraer valores válidos.')
                    continue
                
                # Validar que al menos tengamos nombre o email
                if not values.get('name') and not values.get('email'):
                    errors.append(f'Fila {row_idx}: Debe tener al menos Nombre o Email.')
                    continue
                
                # Agregar categoría según tipo de contacto
                category_name = dict(self._fields['contact_type'].selection).get(self.contact_type)
                if category_name:
                    category = self._get_or_create_category(category_name)
                    values['category_id'] = [(4, category.id)]
                
                # Verificar si hay un conflicto resuelto para esta fila
                if row_idx in conflicts_by_row:
                    conflict = conflicts_by_row[row_idx]
                    
                    if conflict.action == 'skip':
                        # Omitir esta fila
                        continue
                    elif conflict.action == 'update':
                        # Actualizar el contacto existente
                        if conflict.existing_partner_id:
                            update_values = values.copy()
                            update_values.update(import_info)
                            conflict.existing_partner_id.write(update_values)
                            to_update.append(conflict.existing_partner_id.id)
                            updated_partners |= conflict.existing_partner_id
                        else:
                            errors.append(f'Fila {row_idx}: Contacto existente no encontrado.')
                    elif conflict.action == 'create':
                        # Crear nuevo contacto ignorando el existente
                        create_values = values.copy()
                        create_values.update(import_info)
                        new_partner = self.env['res.partner'].create(create_values)
                        to_create.append(new_partner.id)
                        created_partners |= new_partner
                else:
                    # No hay conflicto, crear nuevo contacto
                    create_values = values.copy()
                    create_values.update(import_info)
                    new_partner = self.env['res.partner'].create(create_values)
                    to_create.append(new_partner.id)
                    created_partners |= new_partner
                    
            except Exception as e:
                error_msg = f'Fila {row_idx}: {str(e)}'
                errors.append(error_msg)
                _logger.warning(error_msg)
        
        # Actualizar estadísticas del wizard
        self.total_rows = len(rows)
        self.to_create = len(to_create)
        self.to_update = len(to_update)
        self.with_errors = len(errors)
        self.error_log = '\n'.join(errors) if errors else _('No hay errores.')
        self.created_ids = created_partners
        self.updated_ids = updated_partners
        
        # Crear registro en el historial de importaciones
        self.env['import.history'].create({
            'file_name': self.file_name or 'Archivo sin nombre',
            'file_type': self.file_type,
            'import_date': import_timestamp,
            'contact_type': self.contact_type,
            'region': self.region,
            'total_rows': len(rows),
            'created_count': len(to_create),
            'updated_count': len(to_update),
            'error_count': len(errors),
            'created_ids': [(6, 0, created_partners.ids)],
            'updated_ids': [(6, 0, updated_partners.ids)],
            'error_log': '\n'.join(errors) if errors else False,
            'user_id': self.env.user.id,
        })

    def _extract_values_from_row(self, row, field_mapping):
        """Extrae valores de una fila según el mapeo de campos."""
        values = {}
        
        for col_idx, odoo_field in field_mapping.items():
            if col_idx < len(row):
                value = row[col_idx]
                
                if not value or (isinstance(value, str) and not value.strip()):
                    continue
                
                # Procesar según el tipo de campo
                if odoo_field == 'country_id':
                    # Buscar país por código o nombre
                    country = self.env['res.country'].search([
                        '|', ('code', '=ilike', value),
                        ('name', '=ilike', value)
                    ], limit=1)
                    if country:
                        values[odoo_field] = country.id
                else:
                    values[odoo_field] = value
        
        return values

    def _find_existing_partner(self, values):
        """Busca si un contacto ya existe por documento o email."""
        domain = []
        
        if values.get('vat'):
            domain.append(('vat', '=ilike', values['vat']))
        
        if values.get('email'):
            if domain:
                domain = ['|'] + domain
            domain.append(('email', '=ilike', values['email']))
        
        if not domain:
            return False
        
        return self.env['res.partner'].search(domain, limit=1)

    def _get_or_create_category(self, name):
        """Obtiene o crea una categoría de contacto."""
        category = self.env['res.partner.category'].search([('name', '=', name)], limit=1)
        if not category:
            category = self.env['res.partner.category'].create({'name': name})
        return category

    def action_view_created(self):
        """Abre vista de contactos creados."""
        self.ensure_one()
        return {
            'name': _('Contactos Creados'),
            'type': 'ir.actions.act_window',
            'res_model': 'res.partner',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.created_ids.ids)],
        }

    def action_view_updated(self):
        """Abre vista de contactos actualizados."""
        self.ensure_one()
        return {
            'name': _('Contactos Actualizados'),
            'type': 'ir.actions.act_window',
            'res_model': 'res.partner',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.updated_ids.ids)],
        }
